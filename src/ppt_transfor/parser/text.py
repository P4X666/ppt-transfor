"""文本框解析器：TextFrame → Text 模型。

遍历段落与 run，提取字体/颜色/对齐/间距/行距。
对 placeholder 形状，解析继承自 layout/master 的对齐、字号、颜色。
"""

from __future__ import annotations

from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN

from ppt_transfor.models.schema import Font, Paragraph, Run, Text

# DrawingML 命名空间
NS_A = "http://schemas.openxmlformats.org/drawingml/2006/main"


def _parse_alignment(alignment) -> str | None:
    """对齐方式枚举 → 字符串"""
    if alignment is None:
        return None
    return alignment.name if hasattr(alignment, "name") else str(alignment)


def _parse_alignment_from_para_xml(para) -> str | None:
    """从段落 XML 的 <a:pPr> 读取对齐，作为 para.alignment 的兜底。"""
    try:
        pPr = para._element.find(f"{{{NS_A}}}pPr")
        if pPr is None:
            return None
        algn = pPr.get("algn")
        if not algn:
            return None
        algn_map = {
            "l": "LEFT",
            "ctr": "CENTER",
            "r": "RIGHT",
            "just": "JUSTIFY",
            "dist": "DISTRIBUTE",
            "thaiDist": "THAI_DISTRIBUTE",
            "justLow": "JUSTIFY_LOW",
        }
        return algn_map.get(algn, algn.upper())
    except Exception:
        return None


def _parse_anchor(anchor) -> str | None:
    """垂直对齐枚举 → 字符串"""
    if anchor is None:
        return None
    return anchor.name if hasattr(anchor, "name") else str(anchor)


def _parse_auto_size(auto_size) -> str | None:
    """自适应大小枚举 → 字符串。

    NONE（显式无自适应）与 None（继承）统一返回 None，
    避免往返时因 add_textbox 默认行为差异产生噪声。
    """
    if auto_size is None:
        return None
    name = auto_size.name if hasattr(auto_size, "name") else str(auto_size)
    # NONE 与继承统一，往返一致
    if name == "NONE":
        return None
    return name


def _parse_bodyPr_wrap(tf_element) -> bool | None:
    """从 <a:bodyPr> 的 wrap 属性读取自动换行。

    python-pptx 的 tf.word_wrap 可能未暴露某些继承或默认设置，
    直接读 XML 作为补充。
    """
    if tf_element is None:
        return None
    try:
        bodyPr = tf_element.find(f"{{{NS_A}}}bodyPr")
        if bodyPr is None:
            return None
        wrap_val = bodyPr.get("wrap")
        if wrap_val is None:
            return None
        # square: 自动换行；none: 不换行
        return wrap_val.lower() == "square"
    except Exception:
        return None


def _parse_bodyPr_insets(tf_element) -> dict[str, int]:
    """从 <a:bodyPr> 读取内部边距 lIns/tIns/rIns/bIns（EMU）。"""
    insets = {}
    if tf_element is None:
        return insets
    try:
        bodyPr = tf_element.find(f"{{{NS_A}}}bodyPr")
        if bodyPr is None:
            return insets
        for attr, key in (
            ("lIns", "margin_left"),
            ("tIns", "margin_top"),
            ("rIns", "margin_right"),
            ("bIns", "margin_bottom"),
        ):
            val = bodyPr.get(attr)
            if val is not None:
                try:
                    insets[key] = int(val)
                except (ValueError, TypeError):
                    pass
    except Exception:
        pass
    return insets


def _extract_para_default_props(para, prs=None) -> dict:
    """从段落级 <a:pPr>/<a:defRPr> 提取默认字体属性。

    很多文本框把字号、颜色、字体名定义在段落默认 run 属性里，
    而不是 <a:lstStyle> 或 run 级别，需要单独解析并固化到 run 上。

    Args:
        para: python-pptx Paragraph 对象
        prs: Presentation 对象（用于主题色固化）

    Returns:
        dict: {"font_size": ..., "font_color": Color, "font_name": ..., "font_cap": ...}
    """
    from ppt_transfor.models.schema import Color
    from ppt_transfor.utils.inheritance import _resolve_schemeclr_to_rgb

    props = {"font_size": None, "font_color": None, "font_name": None, "font_cap": None}
    try:
        pPr = para._element.find(f"{{{NS_A}}}pPr")
        if pPr is None:
            return props
        def_rPr = pPr.find(f"{{{NS_A}}}defRPr")
        if def_rPr is None:
            return props

        # 字号
        sz = def_rPr.get("sz")
        if sz is not None:
            try:
                props["font_size"] = int(int(sz) * 127)
            except (ValueError, TypeError):
                pass

        # 大小写
        cap = def_rPr.get("cap")
        if cap:
            props["font_cap"] = cap.lower()

        # 颜色
        solid_fill = def_rPr.find(f"{{{NS_A}}}solidFill")
        if solid_fill is not None:
            srgb = solid_fill.find(f"{{{NS_A}}}srgbClr")
            if srgb is not None:
                val = srgb.get("val")
                if val:
                    props["font_color"] = Color(type="rgb", value=val)
            else:
                scheme = solid_fill.find(f"{{{NS_A}}}schemeClr")
                if scheme is not None:
                    val = scheme.get("val")
                    if val:
                        rgb_value = _resolve_schemeclr_to_rgb(val, prs)
                        if rgb_value is not None:
                            props["font_color"] = Color(type="rgb", value=rgb_value)
                        else:
                            props["font_color"] = Color(type="theme", value=val)

        # 字体名
        latin = def_rPr.find(f"{{{NS_A}}}latin")
        if latin is not None:
            typeface = latin.get("typeface")
            if typeface:
                props["font_name"] = typeface
    except Exception:
        pass

    return props


def _parse_font(font, prs=None, inherited_props: dict | None = None) -> Font:
    """解析 Font 对象。

    对继承的属性（size/color/name 为 None），从 inherited_props 取值固化。

    Args:
        font: python-pptx Font 对象
        prs: Presentation 对象（用于主题色固化）
        inherited_props: placeholder 继承属性 dict
    """
    from ppt_transfor.utils.color import parse_color

    model = Font()

    # 字体名：显式优先，继承兜底
    if font.name is not None:
        model.name = font.name
    elif inherited_props and inherited_props.get("font_name"):
        model.name = inherited_props["font_name"]

    # 字号：显式优先，继承兜底（关键：避免 SHAPE_TO_FIT_TEXT 基于错误字号撑大）
    if font.size is not None:
        model.size = int(font.size)
    elif inherited_props and inherited_props.get("font_size") is not None:
        model.size = inherited_props["font_size"]

    if font.bold is not None:
        model.bold = font.bold
    if font.italic is not None:
        model.italic = font.italic
    if font.underline is not None:
        model.underline = font.underline

    # 大小写：all 表示全大写，small 表示小型大写
    cap_set = False
    try:
        if font.cap is not None:
            cap_name = font.cap.name if hasattr(font.cap, "name") else str(font.cap)
            if cap_name != "NONE":
                model.cap = cap_name.lower()
                cap_set = True
    except Exception:
        # python-pptx 未暴露 cap 属性时，从 XML 读取
        try:
            cap_val = font._element.get("cap")
            if cap_val:
                model.cap = cap_val.lower()
                cap_set = True
        except Exception:
            pass
    if not cap_set and inherited_props and inherited_props.get("font_cap"):
        model.cap = inherited_props["font_cap"]

    # 颜色：显式优先，继承兜底（主题色已固化为 RGB）
    try:
        color = parse_color(font.color, prs)
        if color is not None:
            model.color = color
        elif inherited_props and inherited_props.get("font_color"):
            model.color = inherited_props["font_color"]
    except Exception:
        if inherited_props and inherited_props.get("font_color"):
            model.color = inherited_props["font_color"]

    return model


def parse_paragraph(para, prs=None, inherited_props: dict | None = None, default_alignment: str | None = None) -> Paragraph:
    """解析单个段落。

    Args:
        para: python-pptx 段落对象
        prs: Presentation 对象（用于主题色固化）
        inherited_props: placeholder 继承属性 dict（对齐/字号/颜色/字体名）
        default_alignment: 当所有来源均未取到对齐时的兜底对齐
    """
    model = Paragraph()

    # 对齐：显式优先，段落 XML 兜底，继承兜底，默认兜底
    align = _parse_alignment(para.alignment)
    if align is None:
        align = _parse_alignment_from_para_xml(para)
    if align is not None:
        model.alignment = align
    elif inherited_props and inherited_props.get("alignment"):
        model.alignment = inherited_props["alignment"]
    elif default_alignment is not None:
        model.alignment = default_alignment

    model.level = para.level

    # 段落级默认字体属性（pPr/defRPr），优先级高于文本框级默认
    para_defaults = _extract_para_default_props(para, prs)
    merged_props = dict(inherited_props) if inherited_props else {}
    for key, value in para_defaults.items():
        if value is not None:
            merged_props[key] = value

    # 间距：space_before/space_after 是 Length（EMU）
    if para.space_before is not None:
        model.space_before = int(para.space_before)
    if para.space_after is not None:
        model.space_after = int(para.space_after)

    # 行距：float 为倍数，Length 为精确值（存 EMU 整数）
    if para.line_spacing is not None:
        ls = para.line_spacing
        if isinstance(ls, float):
            model.line_spacing = ls
        else:
            # Length 对象，存 EMU 整数
            model.line_spacing = float(int(ls))

    # runs
    for run in para.runs:
        model.runs.append(Run(text=run.text, font=_parse_font(run.font, prs, merged_props)))

    return model


def parse_text_frame(tf, prs=None, inherited_props: dict | None = None, default_alignment: str | None = None) -> Text:
    """解析 TextFrame → Text 模型。

    Args:
        tf: python-pptx TextFrame 对象
        prs: Presentation 对象（用于主题色固化）
        inherited_props: placeholder 继承属性 dict（对齐/字号/颜色/字体名）
        default_alignment: 当所有来源均未取到对齐时的兜底对齐
    """
    from ppt_transfor.utils.inheritance import extract_txbody_default_props

    model = Text()

    # 自动换行：优先 API，XML 兜底
    if tf.word_wrap is not None:
        model.word_wrap = tf.word_wrap
    else:
        wrap_from_xml = _parse_bodyPr_wrap(tf._element)
        if wrap_from_xml is not None:
            model.word_wrap = wrap_from_xml

    # 自适应大小
    try:
        if tf.auto_size is not None:
            model.auto_size = _parse_auto_size(tf.auto_size)
    except Exception:
        pass

    # 垂直对齐
    try:
        if tf.vertical_anchor is not None:
            model.vertical_anchor = _parse_anchor(tf.vertical_anchor)
    except Exception:
        pass

    # 文本框内部边距
    insets = _parse_bodyPr_insets(tf._element)
    for key, value in insets.items():
        setattr(model, key, value)

    # shape 自身 txBody 的默认文本样式（lstStyle/lvl1pPr/defRPr）
    # 普通文本框的颜色/字号/对齐常定义在这里，而非 run 级别
    shape_defaults = extract_txbody_default_props(tf._element, prs)

    # 合并继承属性：shape 自身默认样式优先，placeholder layout/master 继承兜底
    merged_props = dict(shape_defaults)
    if inherited_props:
        for key, value in inherited_props.items():
            if merged_props.get(key) is None:
                merged_props[key] = value

    # 段落
    for para in tf.paragraphs:
        model.paragraphs.append(parse_paragraph(para, prs, merged_props, default_alignment))

    return model
